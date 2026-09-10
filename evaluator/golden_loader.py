# -*- coding: utf-8 -*-
"""联合加载黄金集 Excel 与 golden_dataset.json。"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

from evaluator.paths import GOLDEN_DIR

EXCEL_NAME = "漏损问答黄金测评集.xlsx"
JSON_NAME = "golden_dataset.json"
V1_EXCEL_NAME = "漏损问答黄金测评集_v1.xlsx"
V1_JSON_NAME = "golden_dataset_v1.json"
EXCEL_CANDIDATES = (
    "漏损问答黄金测评集.xlsx",
    "漏损问答黄金测评集_v3.xlsx",
    "漏损问答黄金测评集_2026-08-17.xlsx",
)
NUMERIC_RESULT_TYPES = {"stat", "detail"}

EXCEL_COLUMNS = {
    "编号": "case_id",
    "角色": "roles_raw",
    "指标类型/场景": "indicator_type",
    "场景大类": "scene_big",
    "测试问题": "question",
    "时间范围(原始)": "time_scope_raw",
    "时间范围(锚定)": "time_scope_anchored",
    "组织范围(原始)": "org_scope_raw",
    "组织范围(锚定)": "org_scope_anchored",
    "计算口径/规则": "caliber",
    "SQL": "sql",
    "标准值": "standard_value",
    "结果类型": "result_type",
    "数据可用性": "availability",
    "明细链接": "detail_link",
    "备注": "notes",
}


def default_golden_dir(override: Optional[str | Path] = None) -> Path:
    if override:
        return Path(override)
    return GOLDEN_DIR


def load_golden_json(golden_dir: Optional[Path] = None, *, name: str = JSON_NAME) -> Dict[str, Any]:
    path = default_golden_dir(golden_dir) / name
    if not path.exists():
        raise FileNotFoundError(f"找不到黄金集 JSON: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def _excel_path(golden_dir: Path, candidates: tuple[str, ...] = EXCEL_CANDIDATES) -> Path:
    for name in candidates:
        path = golden_dir / name
        if path.exists():
            return path
    return golden_dir / EXCEL_NAME


def load_excel_rows(
    golden_dir: Optional[Path] = None, *, candidates: tuple[str, ...] = EXCEL_CANDIDATES
) -> List[Dict[str, Any]]:
    path = _excel_path(default_golden_dir(golden_dir), candidates)
    if not path.exists():
        raise FileNotFoundError(f"找不到黄金集 Excel: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    if "黄金测评集" not in wb.sheetnames:
        raise ValueError("Excel 缺少工作表『黄金测评集』")
    ws = wb["黄金测评集"]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(v).strip() if v is not None else "" for v in next(rows_iter)]
    out: List[Dict[str, Any]] = []
    for raw in rows_iter:
        if not raw or raw[0] is None:
            continue
        item: Dict[str, Any] = {}
        for idx, col in enumerate(header):
            key = EXCEL_COLUMNS.get(col, col)
            item[key] = raw[idx] if idx < len(raw) else None
        item["case_id"] = str(item.get("case_id") or "").strip()
        if item["case_id"]:
            out.append(item)
    wb.close()
    return out


def is_numeric_case(case: Dict[str, Any]) -> bool:
    return str(case.get("result_type") or "") in NUMERIC_RESULT_TYPES


def merge_golden(
    golden_dir: Optional[Path] = None,
    *,
    json_name: str = JSON_NAME,
    excel_candidates: tuple[str, ...] = EXCEL_CANDIDATES,
) -> Dict[str, Any]:
    """Excel 提供题面与 SQL 列，JSON 补充容差、明细结构和历史标准值。"""
    payload = load_golden_json(golden_dir, name=json_name)
    excel_rows = {
        row["case_id"]: row for row in load_excel_rows(golden_dir, candidates=excel_candidates)
    }
    merged = []
    for case in payload.get("cases", []):
        extra = excel_rows.get(case["case_id"], {})
        row = dict(case)
        row["excel"] = extra
        row["numeric_evaluable"] = is_numeric_case(case)
        merged.append(row)
    missing = [cid for cid in excel_rows if cid not in {c["case_id"] for c in merged}]
    return {
        "meta": payload.get("meta") or {},
        "cases": merged,
        "excel_count": len(excel_rows),
        "json_count": len(payload.get("cases", [])),
        "numeric_count": sum(1 for c in merged if c["numeric_evaluable"]),
        "excel_only": missing,
    }


def merge_v1_golden(golden_dir: Optional[Path] = None) -> Dict[str, Any]:
    """从扁平化的 data/golden 文件加载 v1 历史黄金集。"""
    return merge_golden(
        golden_dir,
        json_name=V1_JSON_NAME,
        excel_candidates=(V1_EXCEL_NAME,),
    )


def case_by_id(golden: Dict[str, Any], case_id: str) -> Dict[str, Any]:
    for case in golden["cases"]:
        if case["case_id"] == case_id:
            return case
    raise KeyError(case_id)


def load_agent_answers(name: str, golden_dir: Optional[Path] = None) -> Dict[str, Any]:
    path = default_golden_dir(golden_dir) / name
    return json.loads(path.read_text(encoding="utf-8"))
