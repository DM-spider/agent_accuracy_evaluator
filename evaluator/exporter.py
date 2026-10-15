# -*- coding: utf-8 -*-
"""导出 JSON / Excel / 离线 HTML 报告。"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict

from openpyxl import Workbook

from evaluator.contract_loader import load_contracts
from evaluator.evaluation_metrics import run_evaluation_metrics
from evaluator.repository import Repository


def _case_judgments(repo: Repository, run_id: str) -> Dict[str, Any]:
    return (run_evaluation_metrics(repo, run_id, load_contracts()).get("case_judgments") or {})


def export_json(repo: Repository, run_id: str, dest: Path) -> Path:
    run = repo.get_run(run_id)
    cases = []
    for item in repo.list_cases(run_id):
        detail = repo.load_case_detail(run_id, item["case_id"]) or item
        cases.append(detail)
    payload = {"run": run, "cases": cases}
    dest.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return dest


def export_xlsx(repo: Repository, run_id: str, dest: Path) -> Path:
    run = repo.get_run(run_id) or {}
    summary = run.get("summary") or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "概览"
    ws.append(["指标", "值"])
    for key in (
        "run_id",
        "status",
        "consistency_threshold",
        "case_pass_rate",
        "numeric_accuracy",
        "assessability_rate",
        "numeric_coverage",
        "row_coverage",
        "unexpected_rate",
        "auto_parse_rate",
        "scored_cases",
        "not_scored_cases",
        "pass_cases",
        "fail_cases",
        "review_cases",
    ):
        ws.append([key, summary.get(key, run.get(key))])
    for key, value in (summary.get("timing_summary") or {}).items():
        ws.append([key, value])
    cases_ws = wb.create_sheet("题目")
    judgments = _case_judgments(repo, run_id)
    cases_ws.append(["题号", "状态", "最终结论", "人工平反", "一致性", "场景", "主要失败", "准确率", "覆盖率", "错误值", "缺失", "多余", "轮次", "完成状态", "首段回答ms", "完整响应ms", "SQL耗时ms", "时间筛选", "对齐问题"])
    for item in repo.list_cases(run_id):
        metrics = item.get("metrics") or {}
        judgment = judgments.get(item.get("case_id")) or {}
        cases_ws.append(
            [
                item.get("case_id"),
                item.get("status"),
                judgment.get("final_verdict"),
                judgment.get("manual_verdict"),
                judgment.get("consistency"),
                (repo.load_case_detail(run_id, item["case_id"]) or {}).get("result", {}).get("scene_big"),
                item.get("primary_failure"),
                metrics.get("accuracy") if item.get("status") not in {"REVIEW", "NOT_SCORED"} else None,
                metrics.get("coverage"),
                metrics.get("wrong_count"),
                metrics.get("missing_count"),
                metrics.get("unexpected_count"),
                item.get("turn_index"),
                item.get("completion_status"),
                (item.get("agent_timings") or {}).get("first_answer_ms"),
                (item.get("agent_timings") or {}).get("completed_ms"),
                item.get("sql_latency_ms") or None,
                json.dumps((item.get("alignment") or {}).get("reported_periods", [])),
                ", ".join((item.get("alignment") or {}).get("issues", [])),
            ]
        )
    diff_ws = wb.create_sheet("逐项差异")
    diff_ws.append(["题号", "坐标", "指标", "SQL值", "智能体值", "差值", "状态", "证据"])
    for item in repo.list_cases(run_id):
        detail = repo.load_case_detail(run_id, item["case_id"]) or {}
        result = detail.get("result") or {}
        for cmp in result.get("comparison_items") or []:
            diff_ws.append(
                [
                    item["case_id"],
                    json.dumps(cmp.get("coordinates") or {}, ensure_ascii=False),
                    cmp.get("metric"),
                    cmp.get("expected_value"),
                    cmp.get("actual_value"),
                    cmp.get("delta"),
                    cmp.get("status"),
                    (cmp.get("evidence") or "")[:200],
                ]
            )
    context_ws = wb.create_sheet("期间与证据")
    context_ws.append(["题号", "请求参数", "SQL参数", "期间与口径证据", "Agent原文", "执行SQL", "SQL结果", "平反备注"])
    for item in repo.list_cases(run_id):
        detail = repo.load_case_detail(run_id, item["case_id"]) or {}
        alignment = item.get("alignment") or {}
        sql = detail.get("sql_snapshot") or {}
        context_ws.append([item["case_id"], json.dumps(alignment.get("requested_params", {})),
                           json.dumps(sql.get("params", {})), alignment.get("evidence", ""),
                           (detail.get("agent_answer") or {}).get("text", ""), sql.get("executed_sql", ""),
                           json.dumps(sql.get("rows", []), ensure_ascii=False),
                           (detail.get("review") or {}).get("note", "")])
    recheck_ws = wb.create_sheet("SQL补查")
    recheck_ws.append(["题号", "补查时间", "用途", "参数", "执行SQL", "耗时ms", "错误", "结果行"])
    for item in repo.list_cases(run_id):
        detail = repo.load_case_detail(run_id, item["case_id"]) or {}
        for snapshot in detail.get("sql_rechecks", []):
            for row in snapshot.get("rows") or [None]:
                recheck_ws.append([item["case_id"], snapshot.get("finished_at"), snapshot.get("note"),
                                   json.dumps(snapshot.get("params", {}), ensure_ascii=False), snapshot.get("executed_sql"),
                                   snapshot.get("latency_ms"), snapshot.get("error"), json.dumps(row, ensure_ascii=False)])
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    wb.save(dest)
    return dest


def export_html(repo: Repository, run_id: str, dest: Path) -> Path:
    from evaluator.report_html import build_report

    dest.write_text(build_report(repo, run_id, load_contracts()), encoding="utf-8")
    return dest
