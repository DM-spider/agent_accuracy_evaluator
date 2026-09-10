import json

import httpx
import pytest

from evaluator.agent_client import redact_headers
from evaluator.answer_context import align_answer, params_for_period
from evaluator.contract_loader import CONTRACTS_PATH, load_contracts
from evaluator.extractor import extract_claims, claims_from_sql_rows
from evaluator.models import CaseStatus, RunStatus, SqlSnapshot
from evaluator.orchestrator import Orchestrator
from evaluator.repository import Repository
from evaluator.run_context import build_run_context, to_pg_params, preview_sql
from evaluator.session_guard import SessionGuard
from evaluator.sql_executor import SqlExecutor
from evaluator.stream_client import StreamAgentClient, safe_time_hints, sse_frames


def frame(kind, **kwargs):
    return "data: " + json.dumps(dict(type=kind, **kwargs), ensure_ascii=False) + "\r\n\r\n"


def done(text=""):
    return frame("response.completed", response={"id": "r", "status": "completed", "output": [
        {"role": "assistant", "content": [{"type": "output_text", "text": text}]}]})


class Chunks(httpx.SyncByteStream):
    def __init__(self, content, fail=False):
        self.content, self.fail = content.encode(), fail

    def __iter__(self):
        for i in range(0, len(self.content), 7):
            yield self.content[i:i + 7]
        if self.fail:
            raise httpx.ReadTimeout("secret-cookie")


def client_for(content, fail=False, calls=None):
    def handle(request):
        if calls is not None:
            calls.append(json.loads(request.content))
        assert request.headers["cookie"] == "SESSION_ID=secret-cookie"
        return httpx.Response(200, headers={"content-type": "text/event-stream"}, stream=Chunks(content, fail))
    return StreamAgentClient(session_id="s", task_id="t", template_id="tpl", org_id="org",
                             login_session="secret-cookie", transport=httpx.MockTransport(handle))


@pytest.fixture
def contract():
    return load_contracts(CONTRACTS_PATH)[0]


@pytest.fixture
def ctx():
    return build_run_context("2026-09-07")


def test_fragmented_unicode_queue_tools_and_completion(contract, ctx):
    content = frame("response.pending", session_id="s") + frame("heartbeat", timestamp="1788768196000")
    content += frame("response.created", response={"id": "r", "status": "in_progress"})
    content += frame("response.output_item.done", item={"type": "function_call", "arguments": "sensitive"})
    content += frame("response.output_text.delta", delta="产销差率为 5%", output_index=1)
    content += done()
    calls = []
    answer = client_for(content, calls=calls).ask(contract, ctx)
    assert answer.text == "产销差率为 5%"
    assert answer.completion_status == "completed" and not answer.error
    assert answer.timings["first_answer_ms"] is not None
    assert answer.timings["completed_ms"] >= answer.timings["first_answer_ms"]
    assert answer.timings["server_duration_ms"] is None
    assert set(calls[0]) == {"sessionId", "taskId", "templateId", "taskName", "message"}
    assert "sensitive" not in answer.model_dump_json()
    assert "secret-cookie" not in answer.model_dump_json()


@pytest.mark.parametrize("ending,fail", [("", False), ("data: [DONE]\n\n", False), ("", True), ("data: {bad}\n\n", False)])
def test_no_false_completion_or_retry(contract, ctx, ending, fail):
    calls = []
    client = client_for(frame("response.output_text.delta", delta="部分回答") + ending, fail, calls)
    answer = client.ask(contract, ctx)
    assert answer.error and answer.completion_status == "unknown"
    assert answer.timings["completed_ms"] is None
    assert answer.text == "部分回答"
    assert len(calls) == 1
    assert client.ask(contract, ctx).completion_status == "not_sent"
    assert len(calls) == 1 and answer.retries == 0
    assert "secret-cookie" not in answer.model_dump_json()


@pytest.mark.parametrize("bad", [frame("heartbeat", session_id="other"), frame("response.completed", response={"status": "in_progress"})])
def test_reject_mismatched_session_and_nonterminal(contract, ctx, bad):
    assert client_for(bad).ask(contract, ctx).error


def test_sse_multiline_and_unterminated():
    assert list(sse_frames([": ping", "event: test", "data: a", "data: b", ""])) == [("test", "a\nb")]
    assert list(sse_frames(["data: unfinished"])) == []


def test_final_output_excludes_prior_assistant_narration(contract, ctx):
    event = frame("response.completed", response={"status": "completed", "output": [
        {"role": "assistant", "content": [{"type": "output_text", "text": "planning"}]},
        {"role": "assistant", "content": [{"type": "output_text", "text": "final"}]}]})
    assert client_for(event).ask(contract, ctx).text == "final"


def test_cookie_headers_redacted():
    assert set(redact_headers({"Cookie": "x", "x-session-id": "x", "Authorization": "x"}).values()) == {"***"}


ANSWER = "集团本部 2026年7月（当月）\n| 指标 | 数值 |\n|---|---|\n| **供水量** | 5,038.00 万 m³ |\n| **售水量** | 4,441.19 万 m³ |\n| 产销差率 | 11.85% |\n\n环比\n2026年6月 9.47%\n年累计 2026年7月 7.91%"


def test_period_primary_scope_and_organization(contract, ctx):
    info, block = align_answer(contract, ctx, ANSWER)
    assert info["sql_params"] == {"period": 202607}
    assert info["requested_params"] == {"period": 202609}
    assert info["reported_periods"] == [202607]
    assert info["status"] == "REVIEW"
    assert "ORGANIZATION_MISMATCH" in info["issues"]
    assert "PERIOD_FALLBACK_REVIEW" in info["issues"]
    assert "年累计" not in block
    values = {c.metric: c.value for c in extract_claims(contract, block)}
    assert values == pytest.approx({"供水量": 50380000, "售水量": 44411900, "产销差率": 11.85})
    assert claims_from_sql_rows(contract, [{"supply": 50380000}])[0].value == 50380000


def test_safe_tool_time_hint_without_retaining_payload():
    item = {"type": "function_call", "arguments": json.dumps({"period": 202608, "token": "secret"})}
    assert safe_time_hints(item) == [{"field": "period", "value": "202608"}]


def test_last_six_months_matching_window_is_aligned():
    cx02 = next(c for c in load_contracts(CONTRACTS_PATH) if c.case_id == "CX02")
    ctx = build_run_context("2026-09-08")
    text = "| 月份 | 产销差率 |\n|---|---|\n| 2026-03 | 35.34% |\n| 2026-04 | 6.30% |\n| 2026-05 | 16.65% |\n| 2026-06 | 11.57% |\n| 2026-07 | 13.85% |\n| 2026-08 | 14.48% |"
    info, _ = align_answer(cx02, ctx, text)
    assert info["period_source"] == "agent_answer"
    assert info["sql_params"]
    assert "PERIOD_AMBIGUOUS" not in info["issues"]
    assert sorted(info["sql_params"].values()) == [202603, 202604, 202605, 202606, 202607, 202608]


def test_missing_answer_period_uses_question_contract_as_review(contract, ctx):
    info, _ = align_answer(contract, ctx, "集团产销差率为5%")
    assert info["sql_params"] == {"period": 202609}
    assert info["period_source"] == "question_contract"
    assert info["period_confidence"] == "inferred"
    assert "PERIOD_INFERRED" in info["issues"]


def test_tool_period_can_confirm_answer_without_date(contract, ctx):
    info, _ = align_answer(contract, ctx, "集团当月产销差率为5%", [{"field": "period", "value": "202608"}])
    assert info["sql_params"] == {"period": 202608}
    assert info["period_source"] == "agent_tool"


def test_inferred_period_empty_sql_falls_back_once(tmp_path, contract, ctx):
    calls = []

    class DB(SqlExecutor):
        def query(self, sql_template, params=None):
            calls.append(params)
            if params["period"] == 202609:
                return SqlSnapshot(params=params, columns=["supply", "sales", "rate"],
                                   rows=[{"supply": None, "sales": None, "rate": None}], row_count=1)
            return SqlSnapshot(params=params, columns=["supply", "sales", "rate"],
                               rows=[{"supply": 100, "sales": 90, "rate": 10}], row_count=1)

    repo = Repository(tmp_path / "e.db", tmp_path / "runs")
    orch = Orchestrator(repo, [contract], agent_client=client_for(done("集团当月供水量100，售水量90，产销差率10%")),
                        sql_executor=DB(connect=lambda: None), mode="live", enable_watermark=False)
    orch.start_run(ctx)
    detail = repo.load_case_detail(ctx.run_id, contract.case_id)
    assert calls == [{"period": 202609}, {"period": 202608}]
    assert detail["sql_snapshot"]["params"] == {"period": 202608}
    assert detail["alignment"]["period_source"] == "database_fallback"
    assert [attempt["status"] for attempt in detail["alignment"]["sql_attempts"]] == ["empty", "ok"]


def test_sql_error_does_not_trigger_period_fallback(tmp_path, contract, ctx):
    calls = []

    class DB(SqlExecutor):
        def query(self, sql_template, params=None):
            calls.append(params)
            return SqlSnapshot(params=params, error="database unavailable")

    repo = Repository(tmp_path / "e.db", tmp_path / "runs")
    orch = Orchestrator(repo, [contract], agent_client=client_for(done("集团当月产销差率10%")),
                        sql_executor=DB(connect=lambda: None), mode="live", enable_watermark=False)
    orch.start_run(ctx)
    detail = repo.load_case_detail(ctx.run_id, contract.case_id)
    assert calls == [{"period": 202609}]
    assert detail["sql_snapshot"]["error"] == "database unavailable"
    assert detail["alignment"]["sql_attempts"] == [
        {"params": {"period": 202609}, "status": "error", "row_count": 0}
    ]


def test_period_override_moves_all_month_boundaries(ctx):
    contract = load_contracts(CONTRACTS_PATH)[0].model_copy(update={
        "parameter_resolver": "month_range",
        "sql_template": "SELECT :period, :period_prev, :start_date, :end_date, :next_month_start, :year_start",
    })
    params = params_for_period(contract, ctx, 202608)
    assert params == {"period": 202608, "period_prev": 202607, "start_date": "2026-08-01",
                      "end_date": "2026-08-31", "next_month_start": "2026-09-01", "year_start": "2026-01-01"}


@pytest.mark.parametrize("text", ["2026年6月与2026年7月当月产销差率 5%", "2027年1月当月产销差率 5%"])
def test_ambiguous_period_never_guessed(contract, ctx, text):
    info, _ = align_answer(contract, ctx, text)
    assert info["status"] == "REVIEW" and not info["sql_params"]


def test_relative_month_without_explicit_date_uses_question_context(contract, ctx):
    info, _ = align_answer(contract, ctx, "本月产销差率 5%")
    assert info["sql_params"] == {"period": 202609}
    assert info["period_confidence"] == "inferred"


def test_bind_cast_literal_comment_percent():
    template = "SELECT :period::numeric, ':period', ':period', '12:34', 'a%' -- :ignored\n"
    sql, bound = to_pg_params(template, {"period": 202607})
    assert "%(period)s::numeric" in sql
    assert sql.count("%(__text_period)s") == 2
    assert bound["__text_period"] == "202607"
    assert "'12:34'" in sql and "-- :ignored" in sql
    assert "'202607'" in preview_sql(template, {"period": 202607})
    assert preview_sql("SELECT :x", {"x": "a'b"}) == "SELECT 'a''b'"
    from pg8000.dbapi import convert_paramstyle
    converted, args = convert_paramstyle("pyformat", sql, bound)
    assert "'a%'" in converted
    assert 202607 in args


def test_same_session_ten_sequential_and_independent_results(tmp_path, contract, ctx):
    calls = []
    client = client_for(done(ANSWER), calls=calls)
    contracts = [contract.model_copy(update={"case_id": f"Q{i}"}) for i in range(10)]
    class DB(SqlExecutor):
        def query(self, sql_template, params=None):
            assert params == {"period": 202607}
            return SqlSnapshot(rows=[{"supply": 50380000, "sales": 44411900, "rate": 11.85}], latency_ms=1)
    repo = Repository(tmp_path / "e.db", tmp_path / "runs")
    orch = Orchestrator(repo, contracts, agent_client=client, sql_executor=DB(connect=lambda: None),
                        mode="live", enable_watermark=False, concurrency=10)
    summary = orch.start_run(ctx)
    assert summary.status == RunStatus.COMPLETED
    assert len(calls) == 10 and all(c["sessionId"] == "s" for c in calls)
    assert summary.scored_cases == 0 and summary.review_cases == 10
    assert summary.numeric_accuracy is None
    rows = repo.list_cases(ctx.run_id)
    assert [c["turn_index"] for c in rows] == list(range(1, 11))
    detail = repo.load_case_detail(ctx.run_id, "Q0")
    assert detail["sql_snapshot"]["rows"]
    assert detail["result"]["alignment"]["conditional_numeric_status"] == "PASS"


def test_unknown_completion_stops_remaining_nine(tmp_path, contract, ctx):
    calls = []
    client = client_for(frame("response.output_text.delta", delta="unfinished"), calls=calls)
    repo = Repository(tmp_path / "e.db", tmp_path / "runs")
    contracts = [contract.model_copy(update={"case_id": f"Q{i}"}) for i in range(10)]
    orch = Orchestrator(repo, contracts, agent_client=client, mode="live", enable_watermark=False)
    summary = orch.start_run(ctx)
    assert summary.status == RunStatus.INTERRUPTED and len(calls) == 1
    assert sum(r["not_scored_reason"] == "NOT_SENT_SESSION_BLOCKED" for r in repo.list_cases(ctx.run_id)) == 9


def test_live_missing_db_never_uses_golden(tmp_path, contract, ctx):
    repo = Repository(tmp_path / "e.db", tmp_path / "runs")
    orch = Orchestrator(repo, [contract], golden_cases={contract.case_id: {"expected": {}}},
                        agent_client=client_for(done(ANSWER)), mode="live", enable_watermark=False)
    orch.start_run(ctx)
    detail = repo.load_case_detail(ctx.run_id, contract.case_id)
    assert detail["result"]["not_scored_reason"] == "SQL_FAIL"
    assert detail["agent_answer"]["text"]
    assert detail["sql_snapshot"]["source"] == "live_sql"


def test_unaligned_answer_still_queries_and_compares(tmp_path, contract, ctx):
    calls = []
    class DB(SqlExecutor):
        def query(self, sql_template, params=None):
            calls.append(params)
            return SqlSnapshot(columns=["supply", "sales", "rate"],
                               rows=[{"supply": 123, "sales": 100, "rate": 0.1}], latency_ms=2, row_count=1)
    repo = Repository(tmp_path / "e.db", tmp_path / "runs")
    orch = Orchestrator(repo, [contract], agent_client=client_for(done("本月为8月，累计截至7月")),
                        sql_executor=DB(connect=lambda: None), mode="live", enable_watermark=False)
    summary = orch.start_run(ctx)
    detail = repo.load_case_detail(ctx.run_id, contract.case_id)
    assert len(calls) == 1
    assert calls[0] == detail["alignment"]["requested_params"]
    assert detail["sql_snapshot"]["rows"] == [{"supply": 123, "sales": 100, "rate": 0.1}]
    assert detail["result"]["comparison_items"]
    assert detail["result"]["not_scored_reason"] != "CONTEXT_UNCONFIRMED"
    assert summary.review_cases == 1


def test_agent_fail_still_runs_sql(tmp_path, contract, ctx):
    calls = []
    class DB(SqlExecutor):
        def query(self, sql_template, params=None):
            calls.append(params)
            return SqlSnapshot(columns=["supply", "sales", "rate"],
                               rows=[{"supply": 1, "sales": 1, "rate": 0.1}], row_count=1)
    repo = Repository(tmp_path / "e.db", tmp_path / "runs")
    orch = Orchestrator(repo, [contract], agent_client=client_for(frame("response.output_text.delta", delta="unfinished")),
                        sql_executor=DB(connect=lambda: None), mode="live", enable_watermark=False)
    orch.start_run(ctx)
    detail = repo.load_case_detail(ctx.run_id, contract.case_id)
    assert calls and calls[0]["period"] == 202609
    assert detail["sql_snapshot"]["rows"]
    assert detail["result"]["not_scored_reason"] == "AGENT_FAIL"


def test_empty_current_metrics_fallback_even_if_agent_confirmed(tmp_path, contract, ctx):
    calls = []
    class DB(SqlExecutor):
        def query(self, sql_template, params=None):
            calls.append(params)
            if params["period"] == 202609:
                return SqlSnapshot(params=params, columns=["supply", "sales", "rate"],
                                   rows=[{"supply": None, "sales": None, "rate": None}], row_count=1)
            return SqlSnapshot(params=params, columns=["supply", "sales", "rate"],
                               rows=[{"supply": 100, "sales": 90, "rate": 0.1}], row_count=1)
    repo = Repository(tmp_path / "e.db", tmp_path / "runs")
    orch = Orchestrator(repo, [contract],
                        agent_client=client_for(done("集团 2026年9月当月供水量100，售水量90，产销差率10%")),
                        sql_executor=DB(connect=lambda: None), mode="live", enable_watermark=False)
    orch.start_run(ctx)
    detail = repo.load_case_detail(ctx.run_id, contract.case_id)
    assert [item["period"] for item in calls] == [202609, 202608]
    assert detail["sql_snapshot"]["params"] == {"period": 202608}
    assert detail["alignment"]["period_source"] == "database_fallback"


def test_sql_recheck_preserves_original_evidence(tmp_path, contract, ctx):
    from evaluator.sql_recheck import recheck_sql
    repo = Repository(tmp_path / "e.db", tmp_path / "runs")
    Orchestrator(repo, [contract], agent_client=client_for(done(ANSWER)), mode="live", enable_watermark=False).start_run(ctx)
    before = repo.load_case_detail(ctx.run_id, contract.case_id)
    class DB(SqlExecutor):
        def query(self, template, params=None):
            assert template == contract.sql_template
            return SqlSnapshot(params=params, rows=[{"supply": 123}], columns=["supply"], row_count=1)
    snapshot = recheck_sql(repo, ctx.run_id, contract.case_id, DB(), {"period": 202608})
    after = repo.load_case_detail(ctx.run_id, contract.case_id)
    assert after["result"] == before["result"]
    assert after["agent_answer"] == before["agent_answer"]
    assert after["sql_snapshot"] == before["sql_snapshot"]
    assert after["sql_rechecks"] == [snapshot]
    assert snapshot["purpose"] == "reference_only" and snapshot["finished_at"]
    with pytest.raises(ValueError):
        recheck_sql(repo, ctx.run_id, contract.case_id, DB(), {"sql": "SELECT 1"})


def test_session_guard_survives_new_instance(tmp_path):
    SessionGuard(tmp_path, "s").acquire("run")
    guard = SessionGuard(tmp_path, "s")
    assert guard.blocked()
    with pytest.raises(FileExistsError):
        guard.acquire("other")
    guard.release()
    assert not guard.blocked()


def test_connect_failure_is_not_sent_without_retry(contract, ctx):
    calls = []
    def fail(request):
        calls.append(request)
        raise httpx.ConnectError("connection failed")
    client = StreamAgentClient(session_id="s", task_id="t", template_id="tpl", org_id="org",
                               login_session="cookie", transport=httpx.MockTransport(fail))
    result = client.ask(contract, ctx)
    assert result.completion_status == "not_sent" and client.halted and not client.uncertain
    assert len(calls) == 1


def test_date_range_not_misread_as_first_month(contract, ctx):
    info, _ = align_answer(contract, ctx, "2026年1-5月累计产销差率 5%")
    assert not info["sql_params"]


def test_live_api_missing_credentials_and_invalid_ids(monkeypatch, contract):
    from evaluator import api
    from fastapi import HTTPException
    settings = {"database": {}, "platform": {"enabled": True}}
    monkeypatch.setattr(api, "state", lambda: {"settings": settings, "golden": {}, "contracts": [contract]})
    with pytest.raises(HTTPException) as caught:
        api.create_run(api.CreateRunBody(mode="live"))
    assert caught.value.detail["code"] == "database_not_configured"
    with pytest.raises(HTTPException) as caught:
        api.create_run(api.CreateRunBody(mode="live", anchor_time="2020-01-01"))
    assert caught.value.detail["code"] == "live_anchor_managed"


def test_invalid_completed_shape_stays_an_agent_error(contract, ctx):
    result = client_for(frame("response.completed", response={"status": "completed", "output": ["bad"]})).ask(contract, ctx)
    assert result.error and result.completion_status == "unknown"


def test_remote_error_code_and_message_are_redacted(contract, ctx):
    event = frame("response.failed", response={"id": "r", "status": "failed", "error": {
        "code": "backend_error", "message": "secret-cookie token=private Bearer hidden", "traceback": "internal"}})
    answer = client_for(event).ask(contract, ctx)
    assert answer.completion_status == "failed"
    assert answer.response_body["remote_error"]["code"] == "backend_error"
    assert "private" not in answer.model_dump_json() and "hidden" not in answer.model_dump_json()
    assert "internal" not in answer.model_dump_json()


def test_reports_keep_context_and_escape_answer(tmp_path, contract, ctx):
    from evaluator.exporter import export_html, export_json, export_xlsx
    from openpyxl import load_workbook
    repo = Repository(tmp_path / "e.db", tmp_path / "runs")
    orch = Orchestrator(repo, [contract], agent_client=client_for(done(ANSWER + "\n<script>alert(1)</script>")),
                        mode="live", enable_watermark=False)
    orch.start_run(ctx)
    html = export_html(repo, ctx.run_id, tmp_path / "report.html").read_text(encoding="utf-8")
    assert "&lt;script&gt;" in html and "<script>" not in html
    assert "202607" in html and "未评分" in html
    payload = json.loads(export_json(repo, ctx.run_id, tmp_path / "report.json").read_text(encoding="utf-8"))
    assert payload["cases"][0]["alignment"]["sql_params"] == {"period": 202607}
    workbook = load_workbook(export_xlsx(repo, ctx.run_id, tmp_path / "report.xlsx"))
    assert "期间与证据" in workbook.sheetnames
    assert workbook["题目"]["E2"].value is None
