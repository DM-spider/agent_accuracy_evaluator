# -*- coding: utf-8 -*-
"""build_excel：golden_dataset.json → 漏损问答黄金测评集_<基准日>.xlsx

主 sheet「黄金测评集」列（§12.1）：
编号|角色|指标类型/场景|场景大类|测试问题|时间范围|组织范围
|计算口径/规则|SQL 模板|可执行SQL|标准值|结果类型|数据可用性|备注
明细题另开 sheet「明细_<编号>」，标准值列注明行数。
附页：口径注册表、锚点校验。
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from golden_builder.registry import (RED_RULES, ORG_ROWS, PERIODTYPE_FQC, SRCTYPE_BJ, DIFFRESULT_WLLS,
                      DMA_STATUS, TOLERANCE, ANCHORS, BZ_BM_DICTS, WLLS_MONTH_STALE,
                      FQC_SINGLE_MONTH_NOTE)

MAIN_HEADERS = ["编号", "角色", "指标类型/场景", "场景大类", "测试问题",
                "时间范围", "组织范围",
                "计算口径/规则", "SQL 模板", "可执行SQL", "标准值", "结果类型", "数据可用性", "备注"]

GRAY = PatternFill("solid", fgColor="D9D9D9")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def fmt_stat_fields(case):
    parts = []
    for f in case["expected"].get("fields", []):
        v = f["value"]
        if v is None:
            parts.append(f"{f['label']}=缺失")
        elif f.get("value_scale") == "ratio01" and isinstance(v, (int, float)):
            # 原始值是 0-1 小数，括号内给出换算后的业务单位读数
            parts.append(f"{f['label']}={v:g}（={v * 100:g}{f.get('unit') or '%'}）")
        elif f["unit"] in ("%", "pp"):
            parts.append(f"{f['label']}={v}{f['unit']}")
        else:
            parts.append(f"{f['label']}={v:,.2f} {f['unit']}" if isinstance(v, (int, float)) else f"{f['label']}={v} {f['unit']}")
    return "；".join(parts) if parts else "—"


def build(golden, xlsx_path, B):
    wb = Workbook()
    ws = wb.active
    ws.title = "黄金测评集"
    ws.append(MAIN_HEADERS)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT

    detail_sheets = {}   # case_id -> sheet
    for case in golden["cases"]:
        exp = case.get("expected", {})
        std_val = ""
        if "fields" in exp:
            std_val = fmt_stat_fields(case)
            if "detail" in exp:   # 字段值 + 明细并存（如 GD3/GD4 超时清单）
                n = exp["detail"]["total_rows"]
                std_val += f"；明细见明细_{case['case_id']}（{n} 行）"
        elif "detail" in exp:
            n = exp["detail"]["total_rows"]
            std_val = f"见明细_{case['case_id']}（{n} 行）"
        elif "standard_answer" in exp:
            std_val = exp["standard_answer"]

        row = [case["case_id"], "、".join(case.get("roles", [])), case.get("indicator_type", ""),
               case.get("scene_big", ""), case.get("question", ""),
               case.get("time_scope_raw", "") or case.get("time_scope_anchored", ""),
               case.get("org_scope_raw", "") or case.get("org_scope_anchored", ""),
               case.get("caliber", ""), case.get("sql", ""), case.get("sql_executable", ""), std_val,
               case.get("result_type", ""), case.get("availability", ""), case.get("notes", "")]
        ws.append(row)
        r = ws.max_row
        for c in ws[r]:
            c.alignment = WRAP
        if case.get("availability") in ("no_source", "jump"):
            for c in ws[r]:
                c.fill = GRAY

    # 明细 sheets
    for case in golden["cases"]:
        if "detail" not in case.get("expected", {}):
            continue
        det = case["expected"]["detail"]
        sh = wb.create_sheet(f"明细_{case['case_id']}")
        sh.append([f"{case['case_id']} {case['question']} | 基准日 {B.isoformat()} | "
                   f"共 {det['total_rows']} 行{'（截断至1000）' if det['total_rows'] > 1000 else ''} | "
                   f"口径：{case.get('caliber','')}"])
        sh.cell(row=1, column=1).font = Font(bold=True)
        rows = det["rows"]
        if rows:
            cols = list(rows[0].keys())
            sh.append(cols)
            for c in sh[2]:
                c.fill = HEADER_FILL
                c.font = HEADER_FONT
            for r in rows:
                sh.append([r.get(c) for c in cols])
        detail_sheets[case["case_id"]] = sh

    # 剔除题目说明（业务不理解问题表述）——放在主表之后、明细之前
    if golden.get("meta", {}).get("removed_cases"):
        ws_rm = wb.create_sheet("剔除题目说明", 1)
        ws_rm.append(["题号", "题面", "剔除原因"])
        for c in ws_rm[1]:
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
        for r in golden["meta"]["removed_cases"]:
            ws_rm.append([r["case_id"], r["question"], r["reason"]])
        ws_rm.column_dimensions["A"].width = 10
        ws_rm.column_dimensions["B"].width = 60
        ws_rm.column_dimensions["C"].width = 50

    # 附页：口径注册表
    ws2 = wb.create_sheet("口径注册表")
    ws2.append(["节", "内容"])
    ws2.append(["组织映射（zoneid/level/zonename/bz_bm）", ""])
    for z, lv, n, bm in ORG_ROWS:
        ws2.append(["", f"{z} | level={lv} | {n} | bz_bm={bm}"])
    ws2.append(["periodtype 字典（fqcxfx）", ""])
    for k, v in PERIODTYPE_FQC.items():
        ws2.append(["", f"{k} = {v}"])
    ws2.append(["srctype 字典（bj）", ""])
    for k, v in SRCTYPE_BJ.items():
        ws2.append(["", f"{k} = {v}"])
    ws2.append(["diffresult（wlls）", str(DIFFRESULT_WLLS)])
    ws2.append(["DMA 状态", str(DMA_STATUS)])
    ws2.append(["红字规则 R1~R12", ""])
    for k, v in RED_RULES.items():
        ws2.append(["", f"{k}: {v}"])
    ws2.append(["容差规则", str(TOLERANCE)])
    ws2.append(["数据事实", f"wlls Month 停更(2023-10)={WLLS_MONTH_STALE}; {FQC_SINGLE_MONTH_NOTE}"])
    ws2.append(["各表 bz_bm 实测取值", ""])
    for t, vals in BZ_BM_DICTS.items():
        ws2.append(["", f"{t}: {', '.join(vals)}"])
    for c in ws2[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT

    # 附页：锚点校验（清单，状态见执行报告/verify 输出）
    ws3 = wb.create_sheet("锚点校验")
    ws3.append(["锚点编号", "描述", "期望值", "校验SQL"])
    for a in ANCHORS:
        ws3.append([a["id"], a["desc"], a["expect"], a["sql"]])
    for c in ws3[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT

    # 列宽
    widths = [8, 18, 14, 12, 40, 20, 20, 36, 60, 60, 32, 10, 12, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(xlsx_path)
